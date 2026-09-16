from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent.parent / "public"
OUT.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT / "Matthew_Rundle_Resume.docx"


def set_cellless_page(section):
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)


def set_font(run, name="Arial", size=11, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_bottom_border(paragraph, color="B7C0CC", size="6", space="2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge in ("top", "left", "bottom", "right", "between"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        p_bdr.append(node)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4A5568")
    r_pr.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "19")
    r_pr.append(size)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Resume Section")
    p.add_run(text.upper())
    return p


def add_role(doc, company, location, title, dates):
    p = doc.add_paragraph(style="Role Header")
    left = p.add_run(company)
    set_font(left, size=11.2, bold=True)
    if location:
        loc = p.add_run(f"  |  {location}")
        set_font(loc, size=9.6, color="4A5568")
    p2 = doc.add_paragraph(style="Role Detail")
    title_run = p2.add_run(title)
    set_font(title_run, size=10.2, bold=True)
    date_run = p2.add_run(f"  |  {dates}")
    set_font(date_run, size=9.8, color="4A5568")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Resume Bullet")
    r = p.add_run("- " + text)
    set_font(r)
    return p


def add_compact_line(doc, label, text):
    p = doc.add_paragraph(style="Compact")
    r1 = p.add_run(label)
    set_font(r1, size=9.8, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=9.8)
    return p


doc = Document()
set_cellless_page(doc.sections[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.02

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(1)
title_style.paragraph_format.keep_with_next = True

for name in ["Resume Section", "Role Header", "Role Detail", "Resume Bullet", "Compact"]:
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

section_style = styles["Resume Section"]
section_style.font.name = "Arial"
section_style.font.size = Pt(10.5)
section_style.font.bold = True
section_style.font.color.rgb = RGBColor(0, 0, 0)
section_style.paragraph_format.space_before = Pt(12)
section_style.paragraph_format.space_after = Pt(6)
section_style.paragraph_format.keep_with_next = True

role_header = styles["Role Header"]
role_header.paragraph_format.space_before = Pt(8)
role_header.paragraph_format.space_after = Pt(0)
role_header.paragraph_format.keep_with_next = True

role_detail = styles["Role Detail"]
role_detail.paragraph_format.space_before = Pt(0)
role_detail.paragraph_format.space_after = Pt(2)
role_detail.paragraph_format.keep_with_next = True

bullet_style = styles["Resume Bullet"]
bullet_style.paragraph_format.left_indent = Inches(0.17)
bullet_style.paragraph_format.first_line_indent = Inches(-0.15)
bullet_style.paragraph_format.space_after = Pt(5)
bullet_style.paragraph_format.line_spacing = 1.06
bullet_style.paragraph_format.keep_together = True

compact_style = styles["Compact"]
compact_style.paragraph_format.space_after = Pt(1.3)
compact_style.paragraph_format.line_spacing = 1.0

name = doc.add_paragraph(style="Title")
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.add_run("MATTHEW RUNDLE")
remove_paragraph_borders(name)

headline = doc.add_paragraph()
headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
headline.paragraph_format.space_after = Pt(2)
r = headline.add_run("MANUFACTURING SYSTEMS & APPLIED AI")
set_font(r, size=11.8, bold=True)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(5)
set_font(contact.add_run("Austin, TX  |  "), size=9.5, color="4A5568")
add_hyperlink(contact, "matthewtrundle@gmail.com", "mailto:matthewtrundle@gmail.com")
set_font(contact.add_run("  |  "), size=9.5, color="4A5568")
add_hyperlink(contact, "linkedin.com/in/matthewtrundle", "https://www.linkedin.com/in/matthewtrundle/")
set_font(contact.add_run("  |  "), size=9.5, color="4A5568")
add_hyperlink(contact, "github.com/matthewtrundle", "https://github.com/matthewtrundle")
portfolio = doc.add_paragraph()
portfolio.alignment = WD_ALIGN_PARAGRAPH.CENTER
portfolio.paragraph_format.space_after = Pt(8)
add_hyperlink(portfolio, "Portfolio: matthew-rundle-resume.vercel.app", "https://matthew-rundle-resume.vercel.app/")

summary = doc.add_paragraph()
summary.paragraph_format.space_after = Pt(3)
summary.paragraph_format.line_spacing = 1.02
r = summary.add_run(
    "Systems builder and data leader with 12+ years across manufacturing software, analytics and team "
    "leadership. Built the manufacturing execution system now running Western Magnetics' end-to-end "
    "production process, from procurement through quality, with 100% traceability across every aspect of "
    "the engineering pipeline. Previously founded Domain Labs to build agentic SaaS products and led data "
    "science and commercial analytics at Expedia Group."
)
set_font(r, size=11)

add_section_heading(doc, "Core Expertise")
p = doc.add_paragraph(style="Compact")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Manufacturing Systems  |  Applied AI  |  Agent Workflows  |  Equipment Integration  |  "
    "Data and Analytics Leadership  |  Executive Decision Support"
)
set_font(r, size=9.7)

add_section_heading(doc, "Professional Experience")
add_role(
    doc,
    "Western Magnetics (Westmag)",
    "South San Francisco, CA / Austin, TX",
    "Manufacturing Systems & Applied AI",
    "Jan 2026 - Present",
)
add_bullet(
    doc,
    "Joined as employee #13 at an a16z-backed hard-tech startup shortly after its seed round and helped bring the company out of stealth."
)
add_bullet(
    doc,
    "Built the manufacturing execution system from the ground up; it now runs Western Magnetics' end-to-end manufacturing process, spanning procurement, production, quality and supply, with 100% traceability across every aspect of the engineering pipeline."
)
add_bullet(
    doc,
    "Partnered with cross-functional teams to translate factory needs into connected software; worked with mechatronics engineers to build custom controllers and integrations that pull live signals directly from imported production equipment, often reverse-engineering undocumented vendor protocols."
)
add_bullet(
    doc,
    "Built design and test tooling around motor development, including agentic hill-climb experiment loops run in partnership with frontier model teams."
)
add_bullet(
    doc,
    "Built the company website (with a graphic designer), the NDA platform, and the employee check-in/check-out system."
)

add_role(
    doc,
    "Domain Labs",
    "Austin, TX",
    "Founder, AI and SaaS Solutions",
    "2025",
)
add_bullet(
    doc,
    "Partnered with small businesses to design and build custom SaaS applications, replacing dependencies on Shopify, Squarespace, and fragmented point solutions."
)
add_bullet(
    doc,
    "Built CoreLinq Communications from the ground up: email, SMS, drip campaigns, cold and warm-lead calling, orchestrated through an agentic process with order tracking and fulfillment, coordinating customer outreach and turning conversations into useful next steps."
)
add_bullet(
    doc,
    "Built CoreLinq Scribe for dental practices: tracks patient interactions and generates SOAP notes with ICD codes, CDT codes and procedural notes; includes inventory management with predictive analytics that cross-references upcoming appointments and auto-orders from suppliers."
)

doc.add_page_break()
add_section_heading(doc, "Data and Analytics Leadership")

add_role(
    doc,
    "Expedia Group (Vrbo)",
    "Austin, TX",
    "Senior Manager, Supply and Commercial Data Science and Analytics",
    "Apr 2021 - Jan 2025",
)
add_bullet(
    doc,
    "Led supply and commercial analytics to support market prioritization, property acquisition, and business planning."
)
add_bullet(
    doc,
    "Established an executive analytics function delivering regular insights for commercial strategy and resource allocation decisions."
)
add_bullet(
    doc,
    "Led development of listing quality scoring to support property performance analysis and commercial prioritization."
)
add_bullet(
    doc,
    "Built and managed a team of six data scientists and established shared Python and SQL standards for analytics work."
)
add_bullet(
    doc,
    "Integrated disparate data sources into executive reporting to support consistent business decisions."
)

add_role(
    doc,
    "Expedia Group (Vrbo)",
    "Austin, TX",
    "Manager, Analytics and Data Science",
    "May 2019 - Apr 2021",
)
add_bullet(
    doc,
    "Developed market segmentation across more than 15 variables to identify and prioritize high-value property acquisition targets."
)
add_bullet(
    doc,
    "Built regional reporting and analytics frameworks that improved sales visibility, prioritization, and commercial decision support."
)

add_section_heading(doc, "Earlier Experience")
add_compact_line(doc, "HomeAway and Vrbo  |  ", "Global Reporting and Analytics Manager; Data Analyst II  |  2018 - 2019")
add_compact_line(doc, "National Instruments  |  ", "Senior Consultant, Data Analytics  |  2016 - 2018")
add_compact_line(doc, "The Advisory Board Company  |  ", "Senior Business Analyst  |  2013 - 2016")

add_section_heading(doc, "Technical Skills")
add_compact_line(doc, "Manufacturing Systems  |  ", "MES, traceability, procurement and quality workflows, production data, equipment integration")
add_compact_line(doc, "AI Systems  |  ", "LLM APIs, retrieval-augmented generation, vector databases, agent workflows, natural-language-to-SQL, prompt design")
add_compact_line(doc, "Engineering  |  ", "Python, SQL, TypeScript, JavaScript, Next.js, React, PostgreSQL, Supabase, API integration, custom controllers and signal ingestion, Docker, CI/CD")
add_compact_line(doc, "Data and Cloud  |  ", "Databricks, Tableau, experimentation, segmentation, statistical analysis, AWS, analytics pipelines")
add_compact_line(doc, "Leadership  |  ", "AI strategy, product delivery, executive communication, team building, stakeholder management, operating model design")

add_section_heading(doc, "Education")
p = doc.add_paragraph(style="Compact")
r1 = p.add_run("Washington State University")
set_font(r1, size=10.1, bold=True)
r2 = p.add_run("  |  BA, Management Information Systems  |  2006")
set_font(r2, size=10.1)

for section in doc.sections:
    set_cellless_page(section)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    rr = footer.add_run("Matthew Rundle  |  Manufacturing Systems & Applied AI")
    set_font(rr, size=8, color="6B7280")

doc.core_properties.title = "Matthew Rundle Resume"
doc.core_properties.subject = "Resume for manufacturing systems and applied AI roles"
doc.core_properties.author = "Matthew Rundle"
doc.save(DOCX_PATH)
print(DOCX_PATH.resolve())
